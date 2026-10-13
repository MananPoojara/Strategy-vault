"""Real browser check of admitted/withheld results, reading cards and exports.
Uses one Chromium instance and restores existing account settings afterward.
"""
import asyncio,json,os,zipfile,csv
from openpyxl import load_workbook
from pathlib import Path
from playwright.async_api import async_playwright,expect

CFG=dict(line.split('=',1) for line in (Path.home()/'.cache/strategy-vault-webapp/.env').read_text().splitlines() if line and not line.startswith('#'))
URL=CFG['VAULT_PUBLIC_URL']
OUT=Path.home()/'.cache/strategy-vault-webapp/accuracy-browser-2026-10-13'
OUT.mkdir(exist_ok=True)

async def ready(page,expression):
 for _ in range(300):
  if await page.evaluate('() => ('+expression+')'):return
  await asyncio.sleep(.1)
 raise AssertionError('The requested result context did not finish loading')

async def main():
 async with async_playwright() as p:
  browser=await p.chromium.launch(executable_path='/usr/bin/google-chrome',headless=True,args=['--disable-dev-shm-usage'])
  context=await browser.new_context(viewport={'width':1440,'height':1000},accept_downloads=True)
  page=await context.new_page();errors=[];checks=[];original=None
  page.on('pageerror',lambda e:errors.append(str(e)))
  try:
   await page.goto(URL+'/login.html')
   await page.locator('#email').fill(CFG['VAULT_ADMIN_EMAIL']);await page.locator('#password').fill(CFG['VAULT_ADMIN_PASSWORD'])
   await page.get_by_role('button',name='Sign in',exact=True).click();await page.wait_for_url('**/full-variant1.html**')
   original=(await (await context.request.get(URL+'/api/app/me')).json())['settings']
   await page.goto(URL+'/full-variant1.html#/')
   await expect(page.locator('.gallery > .legend-card')).to_have_count(71,timeout=45000)
   covel=page.locator('.gallery > .legend-card').filter(has=page.get_by_role('heading',name='Michael Covel & the Turtles',exact=True))
   await expect(covel).to_have_count(1)
   assert '–' not in await covel.locator('.lc-stats').inner_text()
   checks.append('All71author/collectioncards readable; Covel has no blank Sharpe statistic')
   await ready(page,"HOME_STATS?.connors?.tested>0 && HOME_STATS?.chande?.tested>0 && BOOK_COVERAGE.authors.connors && !document.querySelector('#flipov')");await asyncio.sleep(.6)
   await page.screenshot(path=str(OUT/'library.png'))
   await page.goto(URL+'/full-variant1.html#/guru/minervini')
   await expect(page.locator('[data-book-coverage]')).to_be_visible(timeout=30000)
   await expect(page.locator('.bookcard.pv')).to_have_count(2)
   assert 'fundamentals' in (await page.locator('.bookcard.pv').first.inner_text()).lower()
   links=page.locator('.bookcard.pv .source-pdfs a');assert await links.count()>=2
   response=await context.request.get(URL+await links.first.get_attribute('href'),headers={'Range':'bytes=0-63'})
   assert response.status==206 and (await response.body()).startswith(b'%PDF-')
   await page.screenshot(path=str(OUT/'reading-author.png'))
   checks.append('Theory/source books have reviewed summaries and authenticated original PDF links')
   await page.goto(URL+'/full-variant1.html#/s/f1a?from=2020-01-01&to=2025-12-31')
   await expect(page.locator('#csv')).to_be_visible(timeout=30000)
   assert await page.evaluate("S.basket==='nifty50' && S.tk.startsWith('basket:nifty50:')")
   checks.append('Stock-only strategy defaults to an actual Nifty50constituent instead of index bars')
   query='basket=nifty50&tk=basket%3Anifty50%3ARELIANCE&from=2020-01-01&to=2025-12-31&cost=2&slippage=3'
   await page.goto(URL+'/full-variant1.html#/s/f1a?'+query)
   await expect(page.locator('#csv')).to_be_visible(timeout=30000)
   assert await page.locator('.metrics .metric > span').all_text_contents()==['Net Profit','Avg net pnl','No. of trades','Win Rate','Expectancy','MaxDD','Sharpe']
   await ready(page,"STRAT?.tk==='basket:nifty50:RELIANCE' && STRAT.d.costs.fees_bps===0 && !document.querySelector('#flipov')")
   values=await page.evaluate("({sum:STRAT.d.stats.net_profit,mean:STRAT.d.stats.avg_net_pnl,n:STRAT.d.stats.trades,compound:STRAT.d.stats.total_return,shown:[...document.querySelectorAll('[data-tile]')].map((e,i)=>e.textContent===STRAT.tiles[i][2](STRAT.tiles[i][1]))})")
   assert abs(values['sum']-values['mean']*values['n'])<.0002 and all(values['shown'])
   assert abs(values['sum']-values['compound'])>.001
   await expect(page.locator('[data-return-definition]')).to_contain_text('sum of closed-trade')
   await expect(page.locator('[data-benchmark-comparison]')).to_contain_text('Compounded model return')
   assert await page.locator('[data-tile]').count()==7
   assert await page.locator('input[name="cost"]').count()==0 and await page.locator('.tags2').count()==0
   assert await page.locator('[data-market-controls] #open-optimization').count()==1 and await page.locator('[data-market-controls] button[form="period"]').count()==1
   assert await page.locator('[data-market-controls] [data-basket]').all_text_contents()==['NSE 500','Nifty 50','Midcap100','All F&O']
   heat=await page.evaluate('({font:getComputedStyle(document.querySelector(".heat2 td")).fontFamily,total:+document.querySelector("[data-annual-average]").textContent,expected:VaultReport.monthly(STRAT.d).annualAverage*100})')
   assert 'Manrope' in heat['font'] and abs(heat['total']-heat['expected'])<=.0051
   checks.append('Market/Index/Basket and run controls share the requested area; fee input and metadata strip removed; styled heatmap averages reconcile')
   checks.append('Actual stock metrics use closed returnSUM, accurate titles and atomic financial values')
   async with page.expect_download() as dl:await page.locator('#trades-xlsx').click()
   download=await dl.value;file=OUT/'actual-stock-trades.xlsx';await download.save_as(str(file))
   with zipfile.ZipFile(file) as z:assert 'xl/workbook.xml' in z.namelist()
   wb=load_workbook(file,data_only=True);sheet=wb['Trades'];rows=list(sheet.values)
   expected=await page.evaluate('VaultReport.tradeRows(STRAT.d,STRAT.tk)')
   assert rows[0]==tuple(await page.evaluate('VaultReport.tradeHeaders')) and len(rows)==len(expected)+1
   for actual,wanted in zip(rows[1:],expected):
    assert actual[0]==wanted[0] and actual[2]==wanted[2] and actual[7] is None and actual[10] is None and actual[12] is None
    assert abs(actual[13]-wanted[13])<1e-12
   assert sheet['N2'].number_format.startswith('0.00%') and sheet['A1'].fill.fgColor.rgb
   metadata=dict(list(wb['Metadata'].values)[1:]);assert metadata['Selected data key']=='basket:nifty50:RELIANCE' and metadata['Execution fees (bps per side)']==0
   assert 'Sources' in wb.sheetnames and 'Parameters' in wb.sheetnames
   checks.append('Actual-stock17column Excel downloads and matches recorded trades, blank money, colors and captured metadata')
   async with page.expect_download() as dl:await page.locator('#csv').click()
   file=OUT/'actual-stock-trades.csv';await (await dl.value).save_as(str(file))
   with file.open(newline='') as stream:csvrows=list(csv.reader(stream))
   assert len(csvrows)==len(expected)+1 and len(csvrows[0])==17 and csvrows[1][0]=='RELIANCE'
   async with page.expect_download() as dl:await page.locator('[data-chart-snapshot="c-eq"]').click()
   file=OUT/'actual-stock-equity.png';await (await dl.value).save_as(str(file));assert file.read_bytes().startswith(b'\x89PNG\r\n\x1a\n')
   async with page.expect_download() as dl:await page.locator('[data-chart-snapshot="c-dd"]').click()
   file=OUT/'actual-stock-drawdown.png';await (await dl.value).save_as(str(file));assert file.read_bytes().startswith(b'\x89PNG\r\n\x1a\n')
   checks.append('Trade CSV, canvas equity PNG and SVG drawdown PNG download through actual buttons')
   await page.screenshot(path=str(OUT/'actual-stock-strategy.png'),full_page=True)
   # Exercise optimizer rendering using a real bounded2rowgrid, not a fabricated API response.
   result=await context.request.get(URL+'/api/in/guru/connors/optimize?id=f1a&ticker=basket%3Anifty50%3ARELIANCE&from=2020-01-01&to=2025-12-31&cost=2&slippage=3&g.hold=4:5:1')
   assert result.status==200
   res=await result.json()
   rendered=await page.evaluate("res=>{const r=rowOf(res,res.rows[0]);const b=bookRow(res);return {sum:r.m.net_profit,model:r.m.total_return,book:b.m.net_profit,stats:statsM(STRAT.d.stats).net_profit};}",res)
   col=res['cols'].index('net_profit');assert rendered['sum']==res['rows'][0][col]
   assert rendered['sum']!=rendered['model'] and rendered['stats']==values['sum']
   checks.append('Optimization row adapters preserve the backend closedSUM rather than overwrite it with compounding')
   for id in ['atm00_fourplay','tdperl_rebo','ss_turtle_soup']:
    await page.goto(URL+'/full-variant1.html#/s/'+id+'?tk=NIFTY+50&from=2000-01-01&to=2025-12-31')
    await expect(page.locator('[data-backtest-state]')).to_be_visible(timeout=30000)
    assert await page.locator('.metrics, #c-eq, #tradesheet').count()==0
   checks.append('Audited extreme/ambiguous methods publish no financial cards or charts')
   await page.set_viewport_size({'width':390,'height':844})
   await page.goto(URL+'/full-variant1.html#/s/f1a?'+query);await expect(page.locator('#csv')).to_be_visible(timeout=30000)
   await page.screenshot(path=str(OUT/'actual-stock-mobile.png'),full_page=True)
   assert not errors,errors
   toolbar=page.get_by_role('toolbar',name='Page actions');box=await toolbar.bounding_box();assert box and box['x']+box['width']<=390.5
   like=page.get_by_role('button',name='Like',exact=True);initial=await like.get_attribute('aria-pressed');await like.click();assert await like.get_attribute('aria-pressed')!=initial
   await page.reload();await expect(page.locator('#csv')).to_be_visible(timeout=30000);await ready(page,"STRAT?.tk==='basket:nifty50:RELIANCE' && !document.querySelector('#flipov')")
   assert await page.get_by_role('button',name='Like',exact=True).get_attribute('aria-pressed')!=initial;await page.get_by_role('button',name='Like',exact=True).click()
   checks.append('Mobile action toolbar fits viewport; real account Like survives reload and restores its prior state; no browser errors')
   # Delay one real response while a newer stock selection finishes. No synthetic prices/results.
   old_seen=asyncio.Event();old_released=asyncio.Event()
   async def delayed(route):
    if 'basket:nifty50:ADANIENT' in __import__('urllib.parse',fromlist=['unquote']).unquote(route.request.url):
     response=await route.fetch();old_seen.set();await asyncio.sleep(5)
     await route.fulfill(response=response);old_released.set()
    else:await route.continue_()
   await page.route('**/api/in/guru/connors/run?**',delayed)
   await page.goto(URL+'/full-variant1.html#/s/f1a?basket=nifty50&tk=basket%3Anifty50%3AADANIENT&from=2021-01-01&to=2025-12-31&cost=7')
   await asyncio.wait_for(old_seen.wait(),20)
   await page.goto(URL+'/full-variant1.html#/s/f1a?'+query)
   await ready(page,"STRAT?.tk==='basket:nifty50:RELIANCE' && !document.querySelector('#flipov')")
   await asyncio.wait_for(old_released.wait(),20);await asyncio.sleep(.5)
   assert await page.evaluate("STRAT.tk==='basket:nifty50:RELIANCE' && STRAT.d.costs.fees_bps===0 && S.tk===STRAT.tk")
   checks.append('A delayed actual prior-stock response cannot replace the current metrics or export context')
   await page.set_viewport_size({'width':1440,'height':1000})
   await page.goto(URL+'/full-variant1.html#/s/f1a?region=us&basket=us_sp500&tk=basket%3Aus_sp500%3AAAPL&from=2000-01-01&to=2020-09-04&slippage=3')
   await expect(page.locator('#csv')).to_be_visible(timeout=30000);await ready(page,"STRAT?.tk==='basket:us_sp500:AAPL' && !document.querySelector('#flipov')")
   assert await page.locator('[data-market-controls] [data-basket]').all_text_contents()==['S&P 500','Nasdaq 100','Dow Jones','Russell 2000']
   assert await page.evaluate("STRAT.d.provenance.trusted_market==='us' && STRAT.d.provenance.currency==='USD' && STRAT.d.costs.fees_bps===0 && STRAT.d.costs.slippage_bps===3")
   await expect(page.locator('[data-basket-note]')).to_contain_text('2020')
   await page.screenshot(path=str(OUT/'us-stock-strategy.png'),full_page=True)
   await page.locator('#basket-universe').select_option('current_plus_past');await ready(page,"STRAT?.tk==='basket:us_sp500_past:AAPL' && !document.querySelector('#flipov')")
   assert await page.evaluate("S.region==='us'&&S.basket==='us_sp500_past'&&STRAT.d.provenance.historical_membership_verified===false")
   await page.goto(URL+'/full-variant1.html#/s/f1a?region=us');await expect(page.locator('#csv')).to_be_visible(timeout=30000);await ready(page,"S.basket==='us_sp500' && STRAT?.tk.startsWith('basket:us_sp500:') && !document.querySelector('#flipov')")
   checks.append('US displays its own four baskets, actual USD stock provenance, source cutoff and both current/past choices; stock rules never default to SPX')

  except Exception:
   await page.screenshot(path=str(OUT/'failure.png'),full_page=True);raise
  finally:
   if original is not None:await context.request.put(URL+'/api/app/settings',data=json.dumps(original),headers={'Origin':URL,'Content-Type':'application/json'})
   await context.close();await browser.close()
   report={'checks':checks,'page_errors':errors};(OUT/'report.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))

if __name__=='__main__':asyncio.run(main())
